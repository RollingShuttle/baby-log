"""
test_rollup.py — the workbook and the day sheet, in isolation. Temp dirs only; nothing here
touches OneDrive or Yisen File.

    python tests/test_rollup.py
"""

# Run straight from the shell, only tests/ is on the path; discovered from the repo root, only
# the root is. Put all three where imports can find them so both ways of running behave alike.
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path[:0] = [str(HERE), str(ROOT), str(ROOT / "tools")]

import json
import shutil
import tempfile
import unittest
from datetime import datetime, timedelta
from unittest import mock

import openpyxl

import rollup
import store

CHILD = "C-20260923-220000-0a1b"
D1 = "2026-09-23"
T1 = "2026-09-23T17:50:00-05:00"
T2 = "2026-09-23T18:10:00-05:00"
T3 = "2026-09-23T20:30:00-05:00"
T4 = "2026-09-23T23:00:00-05:00"

YISEN = {"child_id": CHILD, "name": "Yisen", "born": "2026-09-21",
         "targets": {"feeds_per_day": 8, "wet_per_day": None, "dirty_per_day": None}}


def _rows(ws):
    head = [c.value for c in ws[1]]
    return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


class RollupCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="rollup-test-"))
        self.j = store.Journal(self.tmp / "journal").ensure()
        self.out_dir = self.tmp / "out"
        self.out = self.out_dir / "Baby Log.xlsx"
        self.backups = self.tmp / "backups"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, **kw):
        return rollup.build(self.j, self.out_dir, self.backups, **kw)

    def load(self):
        return openpyxl.load_workbook(self.out, data_only=True)

    def write(self, **kw):
        args = dict(child_id=CHILD, type="feed", time=T1, logged_by="Dad", device="test")
        args.update(kw)
        return self.j.write_event(**args)

    def seed_events(self):
        self.feed = self.write(time=T1, end=T2, note="slow",
                               data={"breast": {"left_s": 300, "right_s": 420, "total_s": 720,
                                                "last_side": "right"},
                                     "bottles": [{"kind": "formula", "ml": 22},
                                                 {"kind": "breast_milk", "ml": 10}],
                                     "made_ml": 60, "leftover_ml": 28})
        self.diaper = self.write(type="diaper", time=T2, logged_by="Mom",
                                 data={"wet": True, "dirty": True, "color": "yellow",
                                       "texture": "seedy", "size": "large", "blowout": True})
        self.sleep = self.write(type="sleep", time=T3, end=T4, data={"where": "crib"})
        self.pump = self.write(type="pump", time=T2, data={"left_ml": 40, "right_ml": 35, "minutes": 15})
        self.growth = self.write(type="growth", time=T3, data={"weight_g": 3420, "length_cm": 50.5})
        self.health = self.write(type="health", time=T3,
                                 data={"medicine": "Vitamin D", "dose": "1 drop", "temp_c": 36.8})
        self.note = self.write(type="note", time=T4, note="First smile", data={"milestone": True})


# -- the Core port ------------------------------------------------------------------------------

class TestCorePort(unittest.TestCase):
    def ev(self, type="feed", time=T1, end=None, data=None, note="", event_id="E-1"):
        return {"event_id": event_id, "type": type, "time": time, "end": end,
                "data": store.validate_data(type, data), "note": note}

    def test_totals_count_on_the_start_day_and_fold_running_timers(self):
        now = datetime.fromisoformat("2026-09-23T18:00:00-05:00")
        evs = [self.ev(data={"breast": {"left_s": 300}, "bottles": [{"kind": "formula", "ml": 22}]}),
               self.ev(time="2026-09-23T17:55:00-05:00",
                       data={"timer": {"side": "left", "side_started": "2026-09-23T17:55:00-05:00"}},
                       event_id="E-2"),
               self.ev(type="diaper", data={"wet": True, "dirty": True}, event_id="E-3"),
               self.ev(type="sleep", time="2026-09-23T23:30:00-05:00", event_id="E-4"),
               self.ev(type="pump", data={"left_ml": 30, "right_ml": 25}, event_id="E-5"),
               self.ev(time="2026-09-24T01:00:00-05:00", event_id="E-6")]
        t = rollup.totals(evs, D1, now)
        self.assertEqual(t, {"feeds": 2, "bottle_ml": 22, "breast_s": 600, "wet": 1, "dirty": 1,
                             "sleeps": 1, "sleep_s": 0, "pumps": 1, "pump_ml": 55})
        later = datetime.fromisoformat("2026-09-24T00:00:00-05:00")
        self.assertEqual(rollup.totals(evs, D1, later)["sleep_s"], 1800)

    def test_breast_seconds_prefers_the_folded_total(self):
        self.assertEqual(rollup.breast_seconds(self.ev(data={"breast": {"left_s": 100, "right_s": 50, "total_s": 900}})), 900)
        self.assertEqual(rollup.breast_seconds(self.ev(data={"breast": {"left_s": 100, "right_s": 50}})), 150)
        self.assertEqual(rollup.breast_seconds(self.ev()), 0)

    def test_last_of_and_usual_gap_and_next_feed(self):
        evs = [self.ev(time="2026-09-23T10:00:00-05:00", event_id="E-1"),
               self.ev(time="2026-09-23T13:00:00-05:00", event_id="E-2"),
               self.ev(time="2026-09-23T15:00:00-05:00", event_id="E-3"),
               self.ev(type="diaper", time="2026-09-23T14:00:00-05:00", event_id="E-4")]
        self.assertEqual(rollup.last_of(evs, "feed", "2026-09-23T14:00:00-05:00")["event_id"], "E-2")
        self.assertEqual(rollup.last_of(evs, "feed", "2026-09-23T16:00:00-05:00")["event_id"], "E-3")
        self.assertIsNone(rollup.last_of(evs, "feed", "2026-09-23T09:00:00-05:00"))
        self.assertIsNone(rollup.usual_gap_ms(evs[:2], "feed"))
        self.assertEqual(rollup.usual_gap_ms(evs, "feed"), 2.5 * 3600 * 1000)
        nxt = rollup.next_feed_at(evs)
        self.assertTrue(store.TIME_RE.match(nxt))
        self.assertEqual(datetime.fromisoformat(nxt).timestamp(),
                         datetime.fromisoformat("2026-09-23T17:30:00-05:00").timestamp())
        self.assertIsNone(rollup.next_feed_at(evs[:2]))

    def test_ties_at_the_same_time_go_to_the_greater_id(self):
        evs = [self.ev(event_id="E-a"), self.ev(event_id="E-b")]
        self.assertEqual(rollup.last_of(evs, "feed", T2)["event_id"], "E-b")

    def test_describe_matches_the_core_examples(self):
        self.assertEqual(rollup.describe(self.ev(data={"breast": {"total_s": 300},
                                                       "bottles": [{"kind": "formula", "ml": 22}]})),
                         "5 min + 22 ml formula")
        self.assertEqual(rollup.describe(self.ev(type="diaper", data={"wet": True, "dirty": True,
                                                                      "color": "yellow"})),
                         "Wet + dirty · yellow")
        self.assertEqual(rollup.describe(self.ev(type="sleep", end="2026-09-23T19:10:00-05:00")),
                         "Sleep 1 h 20 m")
        self.assertEqual(rollup.describe(self.ev(type="pump", data={"left_ml": 30, "right_ml": 30})),
                         "Pumped 60 ml")
        self.assertEqual(rollup.describe(self.ev(type="growth", data={"weight_g": 3420})),
                         "Weight 3.42 kg")
        self.assertEqual(rollup.describe(self.ev(type="health", data={"medicine": "Vitamin D",
                                                                      "dose": "1 drop"})),
                         "Vitamin D 1 drop")
        self.assertEqual(rollup.describe(self.ev(type="note")), "Note")
        self.assertEqual(rollup.describe(self.ev(type="note", data={"milestone": True})), "Milestone")
        self.assertEqual(rollup.describe(self.ev()), "Feeding")
        self.assertEqual(rollup.describe(self.ev(end=T2)), "Feed")
        self.assertEqual(rollup.describe(self.ev(data={"bottles": [{"kind": "formula", "ml": 59}]}), "oz"),
                         "2 oz formula")

    def test_formatting_helpers(self):
        self.assertEqual(rollup.fmt_day(T1), "Wed 23 Sep")
        self.assertEqual(rollup.fmt_day("2026-10-03"), "Sat 3 Oct")
        self.assertEqual(rollup.fmt_time(T1), "17:50")
        self.assertEqual(rollup.day_number("2026-09-21", T1), 3)
        self.assertEqual(rollup.day_number("2026-09-21", "2026-09-21"), 1)
        self.assertEqual(rollup.fmt_amount(22, "ml"), "22 ml")
        self.assertEqual(rollup.fmt_amount(22, "oz"), "0.75 oz")
        self.assertEqual(rollup.fmt_amount(7, "oz"), "0.25 oz")
        self.assertEqual(rollup.fmt_amount(240, "oz"), "8 oz")
        self.assertEqual(rollup.since_text(75 * 60000), "1 h 15 m")
        self.assertEqual(rollup.since_text(45 * 60000), "45 m")
        self.assertEqual(rollup.since_text((2 * 1440 + 180) * 60000), "2 d 3 h")
        self.assertEqual(rollup.since_text(0), "0 m")

    def test_daily_rows_cover_every_date_with_zeros(self):
        evs = [self.ev(time=T1, data={"bottles": [{"kind": "formula", "ml": 22}],
                                      "breast": {"total_s": 90}})]
        rows = rollup.daily_rows(evs, "2026-09-21", "2026-09-23")
        self.assertEqual([r["date"] for r in rows], ["2026-09-21", "2026-09-22", "2026-09-23"])
        self.assertEqual(list(rows[0]), rollup.DAILY_COLUMNS)
        self.assertEqual(rows[0], {"date": "2026-09-21", "feeds": 0, "bottle_ml": 0,
                                   "breast_min": 0.0, "wet": 0, "dirty": 0, "sleeps": 0,
                                   "sleep_min": 0.0, "pumps": 0, "pump_ml": 0})
        self.assertEqual((rows[2]["feeds"], rows[2]["bottle_ml"], rows[2]["breast_min"]), (1, 22, 1.5))
        self.assertEqual(rollup.daily_rows(evs, "bad", "2026-09-23"), [])


# -- the workbook -------------------------------------------------------------------------------

class TestStructure(RollupCase):
    def test_empty_journal_still_produces_a_valid_workbook(self):
        summary = self.build()
        self.assertEqual(Path(summary["path"]), self.out)
        wb = self.load()
        self.assertEqual(wb.sheetnames, rollup.SHEETS + ["About"])
        for name in rollup.SHEETS:
            self.assertEqual(wb[name].max_row, 1, f"{name}: headers only")
        wb.close()
        self.assertEqual(summary["daily"], 0)

    def test_every_sheet_has_the_spec_columns(self):
        self.build()
        wb = self.load()
        expect = {
            "Feeds": "date time end logged_by breast_min left_min right_min bottle_ml formula_ml "
                     "breast_milk_ml made_ml leftover_ml note event_id",
            "Diapers": "date time logged_by wet dirty color texture size rash blowout note event_id",
            "Sleep": "date start end minutes where logged_by note event_id",
            "Daily": "date feeds bottle_ml breast_min wet dirty sleeps sleep_min pumps pump_ml",
        }
        for sheet, cols in expect.items():
            self.assertEqual([c.value for c in wb[sheet][1]], cols.split(), sheet)
        for sheet in ("Pumping", "Growth", "Health", "Notes"):
            head = [c.value for c in wb[sheet][1]]
            self.assertEqual(head[:2], ["date", "time"])
            self.assertEqual(head[-2:], ["note", "event_id"])
            self.assertIn("logged_by", head)
        self.assertEqual([c.value for c in wb["Pumping"][1]], rollup.PUMP_COLUMNS)
        self.assertEqual([c.value for c in wb["Growth"][1]], rollup.GROWTH_COLUMNS)
        self.assertEqual([c.value for c in wb["Health"][1]], rollup.HEALTH_COLUMNS)
        self.assertEqual([c.value for c in wb["Notes"][1]], rollup.NOTE_COLUMNS)
        wb.close()

    def test_about_sheet_warns_that_the_file_is_generated(self):
        self.build()
        wb = self.load()
        text = " ".join(str(c.value) for row in wb["About"].iter_rows() for c in row)
        self.assertIn("do not hand-edit", text)
        self.assertIn(str(self.j.root), text)
        wb.close()

    def test_headers_are_frozen_and_filterable(self):
        self.seed_events()
        self.build()
        wb = self.load()
        self.assertEqual(wb["Feeds"].freeze_panes, "A2")
        self.assertEqual(wb["Feeds"].auto_filter.ref, "A1:N2")
        wb.close()


class TestContent(RollupCase):
    def test_every_type_lands_on_its_sheet_with_a_count(self):
        self.seed_events()
        s = self.build()
        self.assertEqual({k: s[k] for k in ("feeds", "diapers", "sleep", "pumping", "growth",
                                            "health", "notes")},
                         {"feeds": 1, "diapers": 1, "sleep": 1, "pumping": 1, "growth": 1,
                          "health": 1, "notes": 1})
        self.assertGreater(s["bytes"], 0)

    def test_feed_row(self):
        self.seed_events()
        self.build()
        wb = self.load()
        row = _rows(wb["Feeds"])[0]
        wb.close()
        self.assertEqual(row["date"], D1)
        self.assertEqual(row["time"], "17:50")
        self.assertEqual(row["end"], "18:10")
        self.assertEqual(row["logged_by"], "Dad")
        self.assertEqual((row["breast_min"], row["left_min"], row["right_min"]), (12.0, 5.0, 7.0))
        self.assertEqual((row["bottle_ml"], row["formula_ml"], row["breast_milk_ml"]), (32, 22, 10))
        self.assertEqual((row["made_ml"], row["leftover_ml"]), (60, 28))
        self.assertEqual(row["note"], "slow")
        self.assertEqual(row["event_id"], self.feed["event_id"])

    def test_minutes_are_seconds_over_sixty_to_one_place(self):
        self.write(data={"breast": {"total_s": 100}})
        self.build()
        wb = self.load()
        self.assertEqual(_rows(wb["Feeds"])[0]["breast_min"], 1.7)
        wb.close()

    def test_diaper_row(self):
        self.seed_events()
        self.build()
        wb = self.load()
        row = _rows(wb["Diapers"])[0]
        wb.close()
        self.assertEqual((row["time"], row["logged_by"]), ("18:10", "Mom"))
        self.assertEqual((row["wet"], row["dirty"], row["rash"], row["blowout"]), ("yes", "yes", "no", "yes"))
        self.assertEqual((row["color"], row["texture"], row["size"]), ("yellow", "seedy", "large"))
        self.assertEqual(row["event_id"], self.diaper["event_id"])

    def test_sleep_row_has_minutes(self):
        self.seed_events()
        self.build()
        wb = self.load()
        row = _rows(wb["Sleep"])[0]
        wb.close()
        self.assertEqual((row["start"], row["end"], row["minutes"], row["where"]),
                         ("20:30", "23:00", 150.0, "crib"))

    def test_an_end_on_another_day_shows_its_date(self):
        self.write(type="sleep", time=T4, end="2026-09-24T02:30:00-05:00")
        self.build()
        wb = self.load()
        self.assertEqual(_rows(wb["Sleep"])[0]["end"], "2026-09-24 02:30")
        wb.close()

    def test_other_sheets(self):
        self.seed_events()
        self.build()
        wb = self.load()
        pump = _rows(wb["Pumping"])[0]
        growth = _rows(wb["Growth"])[0]
        health = _rows(wb["Health"])[0]
        note = _rows(wb["Notes"])[0]
        wb.close()
        self.assertEqual((pump["left_ml"], pump["right_ml"], pump["total_ml"], pump["minutes"]), (40, 35, 75, 15))
        self.assertEqual((growth["weight_g"], growth["length_cm"], growth["head_cm"]), (3420, 50.5, None))
        self.assertEqual((health["medicine"], health["dose"], health["temp_c"]), ("Vitamin D", "1 drop", 36.8))
        self.assertEqual((note["milestone"], note["note"]), ("yes", "First smile"))

    def test_rows_are_oldest_first(self):
        self.write(time=T3)
        self.write(time=T1)
        self.write(time="2026-09-22T09:00:00-05:00")
        self.build()
        wb = self.load()
        self.assertEqual([r["time"] for r in _rows(wb["Feeds"])], ["09:00", "17:50", "20:30"])
        self.assertEqual([r["date"] for r in _rows(wb["Feeds"])], ["2026-09-22", D1, D1])
        wb.close()

    def test_a_tombstoned_event_leaves_the_workbook(self):
        self.seed_events()
        self.j.delete_event(self.feed["event_id"], reason="test", device="test", edited_by="Dad")
        s = self.build()
        self.assertEqual(s["feeds"], 0)
        self.assertEqual(s["diapers"], 1)

    def test_a_revision_shows_once_with_the_latest_values(self):
        self.seed_events()
        self.j.write_event(event_id=self.feed["event_id"], child_id=CHILD, type="feed", time=T1,
                           data={"bottles": [{"kind": "formula", "ml": 45}]}, device="test")
        self.build()
        wb = self.load()
        rows = _rows(wb["Feeds"])
        wb.close()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["bottle_ml"], 45)
        self.assertEqual(rows[0]["logged_by"], "Dad", "carried through the revision")

    def test_a_running_feed_folds_now_into_breast_min(self):
        started = store.iso_local(datetime.now().astimezone() - timedelta(minutes=10))
        self.write(time=started, data={"breast": {"left_s": 60},
                                       "timer": {"side": "right", "side_started": started}})
        self.build()
        wb = self.load()
        row = _rows(wb["Feeds"])[0]
        wb.close()
        self.assertIsNone(row["end"])
        self.assertGreaterEqual(row["breast_min"], 11.0)
        self.assertLess(row["breast_min"], 12.0)

    def test_cjk_notes_survive(self):
        self.write(note="今天喝得很好", logged_by="妈妈")
        self.build()
        wb = self.load()
        row = _rows(wb["Feeds"])[0]
        wb.close()
        self.assertEqual((row["note"], row["logged_by"]), ("今天喝得很好", "妈妈"))


class TestDaily(RollupCase):
    def test_one_row_per_date_from_the_first_event_to_today(self):
        first = (datetime.now().astimezone() - timedelta(days=3)).strftime("%Y-%m-%d")
        self.write(time=f"{first}T08:00:00-05:00", data={"bottles": [{"kind": "formula", "ml": 22}]})
        self.write(time=f"{first}T12:00:00-05:00", data={"breast": {"total_s": 600}})
        self.write(type="diaper", time=f"{first}T13:00:00-05:00", data={"wet": True})
        s = self.build()
        self.assertEqual(s["daily"], 4)
        wb = self.load()
        rows = _rows(wb["Daily"])
        wb.close()
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0]["date"], first)
        self.assertEqual(rows[-1]["date"], datetime.now().astimezone().strftime("%Y-%m-%d"))
        self.assertEqual((rows[0]["feeds"], rows[0]["bottle_ml"], rows[0]["breast_min"],
                          rows[0]["wet"], rows[0]["dirty"]), (2, 22, 10.0, 1, 0))
        for r in rows[1:]:
            self.assertEqual([r[c] for c in rollup.DAILY_COLUMNS[1:]], [0, 0, 0, 0, 0, 0, 0, 0, 0], r)

    def test_zero_rows_are_values_not_formulas(self):
        self.write(time=T1)
        self.build()
        wb = openpyxl.load_workbook(self.out)         # not data_only: a formula would show as text
        for r in wb["Daily"].iter_rows(min_row=2):
            for c in r[1:]:
                self.assertIsInstance(c.value, (int, float))
        wb.close()

    def test_sleep_minutes_land_on_the_start_day(self):
        self.write(type="sleep", time=T4, end="2026-09-24T02:00:00-05:00")
        self.build()
        wb = self.load()
        rows = {r["date"]: r for r in _rows(wb["Daily"])}
        wb.close()
        self.assertEqual((rows[D1]["sleeps"], rows[D1]["sleep_min"]), (1, 180.0))
        if "2026-09-24" in rows:
            self.assertEqual(rows["2026-09-24"]["sleeps"], 0)


class TestSafety(RollupCase):
    def test_refuses_to_write_while_excel_has_the_file_open(self):
        self.seed_events()
        self.build()
        lock = self.out.with_name(f"~${self.out.name}")
        lock.write_bytes(b"")
        self.assertTrue(rollup.is_locked(self.out_dir))
        with self.assertRaises(RuntimeError) as ctx:
            self.build()
        self.assertIn("open in Excel", str(ctx.exception))
        self.assertEqual(list(self.backups.glob("*.xlsx")), [], "refused before the backup")

    def test_is_locked_is_false_without_the_lock_file(self):
        self.assertFalse(rollup.is_locked(self.out_dir))
        self.build()
        self.assertFalse(rollup.is_locked(self.out_dir))

    def test_backup_is_taken_before_each_overwrite_and_pruned(self):
        self.seed_events()
        self.build()                                     # nothing to back up yet
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 0)
        for _ in range(3):
            self.build()
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 3)
        self.build(keep=2)
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 2)

    def test_two_builds_in_one_second_make_two_backups(self):
        self.build()
        self.build()
        self.build()
        names = sorted(p.name for p in self.backups.glob("Baby Log_*.xlsx"))
        self.assertEqual(len(names), 2)
        self.assertEqual(len(set(names)), 2)
        for n in names:
            self.assertRegex(n, r"^Baby Log_\d{8}-\d{6}-\d{6}\.xlsx$")

    def test_no_temp_debris_left_behind(self):
        self.seed_events()
        self.build()
        self.assertEqual(list(self.out_dir.glob(".tmp-*")), [])
        self.assertEqual(sorted(p.name for p in self.out_dir.iterdir()), ["Baby Log.xlsx"])

    def test_a_failed_save_leaves_no_temp_file(self):
        self.seed_events()
        with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk")):
            with self.assertRaises(OSError):
                self.build()
        self.assertEqual(list(self.out_dir.glob(".tmp-*")), [])
        self.assertFalse(self.out.exists())

    def test_regeneration_is_stable(self):
        self.seed_events()
        first = self.build()
        self.out.unlink()
        second = self.build()
        for k in ("feeds", "diapers", "sleep", "daily"):
            self.assertEqual(first[k], second[k])

    def test_journal_is_not_modified_by_generating_the_rollup(self):
        self.seed_events()
        before = {p.name: p.read_bytes() for p in (self.j.root / "events").rglob("*.json")}
        self.build()
        after = {p.name: p.read_bytes() for p in (self.j.root / "events").rglob("*.json")}
        self.assertEqual(before, after, "the rollup must never write back to the journal")

    def test_an_unreadable_journal_file_is_skipped_with_a_log_line(self):
        self.seed_events()
        day = self.j.root / "events" / "2026-09-23"
        day.mkdir(parents=True, exist_ok=True)
        (day / "E-20260923-000000-0000-r1-abcd.json").write_text("{half", encoding="utf-8")
        self.j.load(force=True)
        with self.assertLogs("baby_log.rollup", level="WARNING") as logs:
            s = self.build()
        self.assertEqual(s["feeds"], 1)
        self.assertTrue(any("E-20260923-000000-0000-r1-abcd.json" in line for line in logs.output))

    def test_output_folder_is_created(self):
        self.assertFalse(self.out_dir.exists())
        self.build()
        self.assertTrue(self.out.exists())


# -- the day sheet ------------------------------------------------------------------------------

class TestDaySheet(RollupCase):
    def html(self, child=YISEN, date=D1, settings=None):
        return rollup.day_sheet_html(child, self.j.events(), date, settings or {"units": "ml"})

    def test_heading_names_the_child_the_day_and_the_day_number(self):
        h = self.html()
        self.assertIn("<h1>Yisen — Feeding &amp; Diapering — Wed 23 Sep (day 3)</h1>", h)
        self.assertIn("<title>Yisen — 2026-09-23</title>", h)

    def test_both_tables_and_their_headers_are_always_present(self):
        h = self.html()
        self.assertIn("<h2>Feeding</h2>", h)
        self.assertIn("<h2>Diapers</h2>", h)
        for col in ("Time", "Breast", "Bottle", "Made / Leftover", "By", "Note", "Wet", "Dirty",
                    "Colour / texture"):
            self.assertIn(f"<th>{col}</th>", h)
        self.assertIn("No feeds", h)
        self.assertIn("No diapers", h)
        self.assertNotIn("<h2>Other</h2>", h)

    def test_feed_and_diaper_rows(self):
        self.seed_events()
        h = self.html()
        self.assertIn("<td class=\"time\">17:50</td>", h)
        self.assertIn("12 min (L 5/R 7)", h)
        self.assertIn("22 ml formula + 10 ml breast milk", h)
        self.assertIn("60 ml / 28 ml", h)
        self.assertIn("<td>Dad</td>", h)
        self.assertIn("<td>slow</td>", h)
        self.assertIn("<td class=\"mark\">✓</td><td class=\"mark\">✓</td>", h)
        self.assertIn("yellow, seedy, large, blowout", h)
        self.assertIn("<td>Mom</td>", h)

    def test_approx_breast_time_is_marked(self):
        self.write(data={"breast": {"total_s": 900, "approx": True}})
        self.assertIn("~15 min", self.html())

    def test_other_table_lists_the_rest_in_time_order(self):
        self.seed_events()
        h = self.html()
        self.assertIn("<h2>Other</h2>", h)
        other = h[h.index("<h2>Other</h2>"):]
        for text in ("Pumped 75 ml", "Sleep 2 h 30 m", "Weight 3.42 kg · Length 50.5 cm",
                     "Vitamin D 1 drop · 36.8 °C", "Milestone", "First smile"):
            self.assertIn(text, other)
        self.assertLess(other.index("Pumped"), other.index("Sleep"))
        self.assertLess(other.index("Sleep"), other.index("Milestone"))
        for kind in ("Pump", "Sleep", "Growth", "Health", "Note"):
            self.assertIn(f"<td>{kind}</td>", other)

    def test_footer_totals_with_targets(self):
        self.seed_events()
        self.write(type="diaper", time=T3, data={"wet": True})
        h = self.html()
        self.assertIn("1 (target 8) feeds · 32 ml bottle · 12 min breast · 2 wet · 1 dirty", h)
        self.assertIn("1 sleeps · 2 h 30 m", h)
        self.assertIn("pumped 75 ml", h)

    def test_footer_without_targets(self):
        self.seed_events()
        child = dict(YISEN, targets={"feeds_per_day": None, "wet_per_day": None, "dirty_per_day": None})
        self.assertIn("1 feeds · 32 ml bottle · 12 min breast · 1 wet · 1 dirty", self.html(child))
        self.assertNotIn("target", self.html(child))

    def test_oz_setting_changes_only_the_formatting(self):
        self.seed_events()
        h = self.html(settings={"units": "oz"})
        self.assertIn("0.75 oz formula", h)
        self.assertIn("2 oz / 1 oz", h)
        self.assertIn("1 oz bottle", h)

    def test_print_css_is_inlined_and_nothing_is_fetched(self):
        h = self.html()
        css = (ROOT / "static" / "print.css").read_text(encoding="utf-8")
        self.assertIn(css, h)
        self.assertIn("@page { size: Letter; margin: 0.5in; }", h)
        self.assertNotIn("http://", h)
        self.assertNotIn("https://", h)
        self.assertNotIn("<link", h)
        self.assertNotIn("<script", h)
        self.assertNotIn("<img", h)

    def test_the_css_is_black_on_white_letter(self):
        css = (ROOT / "static" / "print.css").read_text(encoding="utf-8")
        self.assertIn("size: Letter", css)
        self.assertIn("margin: 0.5in", css)
        self.assertIn("background: #fff", css)
        self.assertIn("color: #000", css)
        self.assertIn(".columns", css)
        self.assertNotIn("url(", css)

    def test_notes_are_escaped(self):
        self.write(note="<b>bold</b> & 22ml")
        h = self.html()
        self.assertIn("&lt;b&gt;bold&lt;/b&gt; &amp; 22ml", h)
        self.assertNotIn("<b>bold</b>", h)

    def test_other_days_are_left_out(self):
        self.write(time="2026-09-22T09:00:00-05:00", note="yesterday")
        self.write(time=T1, note="today")
        h = self.html()
        self.assertIn("today", h)
        self.assertNotIn("yesterday", h)

    def test_without_a_child_it_still_renders(self):
        h = self.html(child=None)
        self.assertIn("<h1>Baby — Feeding &amp; Diapering — Wed 23 Sep</h1>", h)

    def test_missing_print_css_falls_back_with_a_warning(self):
        with mock.patch.object(rollup, "PRINT_CSS", self.tmp / "nope.css"):
            with self.assertLogs("baby_log.rollup", level="WARNING"):
                h = self.html()
        self.assertIn("@page { size: Letter; margin: 0.5in }", h)


class TestWriteDaySheet(RollupCase):
    def test_writes_into_day_sheets_and_matches_the_renderer(self):
        self.seed_events()
        path = rollup.write_day_sheet(self.j, self.out_dir, D1, {"units": "ml"}, child=YISEN)
        self.assertEqual(path, self.out_dir / "Day sheets" / "2026-09-23.html")
        self.assertEqual(path.read_text(encoding="utf-8"),
                         rollup.day_sheet_html(YISEN, self.j.events(), D1, {"units": "ml"}))
        self.assertEqual(list((self.out_dir / "Day sheets").glob(".tmp-*")), [])

    def test_uses_the_journal_child_when_none_is_given(self):
        self.j.write_child(name="Yisen", born="2026-09-21")
        path = rollup.write_day_sheet(self.j, self.out_dir, D1, {"units": "ml"})
        self.assertIn("Yisen — Feeding", path.read_text(encoding="utf-8"))

    def test_overwrites_an_earlier_sheet(self):
        rollup.write_day_sheet(self.j, self.out_dir, D1, {}, child=YISEN)
        self.write(note="new entry")
        path = rollup.write_day_sheet(self.j, self.out_dir, D1, {}, child=YISEN)
        self.assertIn("new entry", path.read_text(encoding="utf-8"))
        self.assertEqual(len(list((self.out_dir / "Day sheets").iterdir())), 1)


class TestRangeHtml(RollupCase):
    def test_one_table_with_the_daily_columns(self):
        self.write(time=T1, data={"bottles": [{"kind": "formula", "ml": 22}]})
        rows = rollup.daily_rows(self.j.events(), "2026-09-22", D1)
        h = rollup.range_html(YISEN, rows, "2026-09-22", D1, {"units": "ml"})
        self.assertIn("<h1>Yisen — Daily totals — Tue 22 Sep to Wed 23 Sep</h1>", h)
        self.assertEqual(h.count("<table"), 1)
        for col in rollup.DAILY_COLUMNS:
            self.assertIn(f">{col}</th>", h)
        self.assertEqual(h.count("<tr>"), 3)
        self.assertIn("<td>2026-09-22</td><td class=\"num\">0</td>", h)
        self.assertIn("<td>2026-09-23</td><td class=\"num\">1</td><td class=\"num\">22</td>", h)
        self.assertNotIn("http", h)
        self.assertIn("@page { size: Letter", h)

    def test_oz_formats_the_volumes(self):
        self.write(time=T1, data={"bottles": [{"kind": "formula", "ml": 59}]})
        rows = rollup.daily_rows(self.j.events(), D1, D1)
        h = rollup.range_html(YISEN, rows, D1, D1, {"units": "oz"})
        self.assertIn("<td class=\"num\">2 oz</td>", h)

    def test_empty_range_says_so(self):
        h = rollup.range_html(YISEN, [], D1, D1, {})
        self.assertIn("No entries", h)


class TestMain(RollupCase):
    def test_main_reads_the_config_and_builds(self):
        cfg = self.tmp / "config.yaml"
        cfg.write_text(json.dumps({"paths": {"app_folder": str(self.j.root),
                                             "output_folder": str(self.out_dir),
                                             "backups": str(self.backups),
                                             "backup_keep": 5}}), encoding="utf-8")
        self.write()
        self.assertEqual(rollup.main(["--config", str(cfg)]), 0)
        self.assertTrue(self.out.exists())


if __name__ == "__main__":
    unittest.main(verbosity=1)
