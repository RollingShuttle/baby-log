"""
rollup.py — the readable files in Yisen File: `Baby Log.xlsx` and the day-sheet HTML.

Both are regenerated whole from the journal and never hand-edited: the journal is the source of
truth, so a manual change here is lost on the next run and deleting either file is harmless. The
workbook is our own file (no images, no rich values), so openpyxl may write it freely; it is
backed up with a microsecond stamp before every overwrite and refused while Excel holds it open
(`~$Baby Log.xlsx` present), because writing then would give OneDrive a conflict copy.

This module also holds the Python port of the `Core` arithmetic the print view and the workbook
need (totals, since-last, the one-line description). It mirrors docs/core.js function for
function so the PC's sheet and the phone's screen never disagree about a number. See SPEC.md §9.
"""
from __future__ import annotations

import html
import logging
import os
import shutil
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import store as store_mod

log = logging.getLogger("baby_log.rollup")

WORKBOOK = "Baby Log.xlsx"
DAY_SHEETS = "Day sheets"

OZ_ML = 29.5735
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# A neutral header: this file is read by tired parents in Excel, not styled for a brand.
HEAD_FILL = PatternFill("solid", fgColor="E7ECEA")
HEAD_FONT = Font(bold=True, size=10)

# The §9 columns, spelled once so the tests and the sheet cannot drift.
FEED_COLUMNS = ["date", "time", "end", "logged_by", "breast_min", "left_min", "right_min",
                "bottle_ml", "formula_ml", "breast_milk_ml", "made_ml", "leftover_ml", "note",
                "event_id"]
DIAPER_COLUMNS = ["date", "time", "logged_by", "wet", "dirty", "color", "texture", "size", "rash",
                  "blowout", "note", "event_id"]
SLEEP_COLUMNS = ["date", "start", "end", "minutes", "where", "logged_by", "note", "event_id"]
PUMP_COLUMNS = ["date", "time", "logged_by", "left_ml", "right_ml", "total_ml", "minutes", "note",
                "event_id"]
GROWTH_COLUMNS = ["date", "time", "logged_by", "weight_g", "length_cm", "head_cm", "note",
                  "event_id"]
HEALTH_COLUMNS = ["date", "time", "logged_by", "temp_c", "medicine", "dose", "symptom", "note",
                  "event_id"]
NOTE_COLUMNS = ["date", "time", "logged_by", "milestone", "note", "event_id"]
DAILY_COLUMNS = ["date", "feeds", "bottle_ml", "breast_min", "wet", "dirty", "sleeps",
                 "sleep_min", "pumps", "pump_ml"]
SHEETS = ["Feeds", "Diapers", "Sleep", "Pumping", "Growth", "Health", "Notes", "Daily"]


def _resource_dir():
    """Where the bundled read-only files live: the PyInstaller temp folder when packaged, beside
    this file from source. Same rule as app.py's; kept here so rollup never imports app."""
    base = getattr(sys, "_MEIPASS", None)
    return Path(base) if base else Path(__file__).resolve().parent


PRINT_CSS = _resource_dir() / "static" / "print.css"


# -- the Core port (docs/core.js §5) -----------------------------------------------------------

def to_dt(x):
    """An aware datetime for an ISO string, a datetime, or None (= now)."""
    if x is None:
        return datetime.now().astimezone()
    if isinstance(x, datetime):
        return x if x.tzinfo else x.astimezone()
    try:
        return datetime.fromisoformat(str(x))
    except ValueError:
        return None


def to_ms(x):
    d = to_dt(x)
    return None if d is None else d.timestamp() * 1000.0


def local_date(iso):
    return str(iso or "")[:10]


def today_local(now=None):
    return to_dt(now).strftime("%Y-%m-%d")


def on_day(events, day):
    return [e for e in events if local_date(e.get("time")) == day]


def is_running(ev):
    """A timer is going. "No end" alone is not enough: the paper import has no end times, and
    those feeds must not read as running for ever. Mirrors Core.isRunning."""
    return (bool(ev) and ev.get("end") is None and ev.get("type") in ("feed", "sleep")
            and bool((ev.get("data") or {}).get("timer")))


def last_of(events, type, before=None):
    """The latest event of a type that started at or before `before`; a running one counts."""
    cut = to_ms(before)
    best, best_ms = None, None
    for ev in events:
        if ev.get("type") != type:
            continue
        ms = to_ms(ev.get("time"))
        if ms is None or ms > cut:
            continue
        if best is None or ms > best_ms or (ms == best_ms and
                                            str(ev.get("event_id")) > str(best.get("event_id"))):
            best, best_ms = ev, ms
    return best


def _median(nums):
    s = sorted(nums)
    if not s:
        return 0
    mid = len(s) // 2
    return s[mid] if len(s) % 2 else (s[mid - 1] + s[mid]) / 2


def usual_gap_ms(events, type, n=6):
    """The median gap between the last n starts of a type; None until there are three."""
    starts = sorted(ms for ms in (to_ms(e.get("time")) for e in events if e.get("type") == type)
                    if ms is not None)[-n:]
    if len(starts) < 3:
        return None
    return _median([b - a for a, b in zip(starts, starts[1:])])


def next_feed_at(events):
    """last feed + usual gap, as a §3.2 string in this machine's zone; a hint, never an alarm."""
    last = last_of(events, "feed")
    gap = usual_gap_ms(events, "feed")
    if last is None or gap is None:
        return None
    return store_mod.iso_local((to_dt(last["time"]) + timedelta(milliseconds=gap)).astimezone())


def bottle_ml(ev):
    return sum((b.get("ml") or 0) for b in ((ev.get("data") or {}).get("bottles") or []))


def breast_seconds(ev, now=None):
    """The folded total (or the sides) plus the running side's elapsed time, floored."""
    br = (ev.get("data") or {}).get("breast") or {}
    s = br["total_s"] if br.get("total_s") is not None else (br.get("left_s") or 0) + (br.get("right_s") or 0)
    timer = (ev.get("data") or {}).get("timer")
    if is_running(ev) and timer and timer.get("side_started"):
        started = to_ms(timer["side_started"])
        if started is not None:
            s += max(0.0, (to_ms(now) - started) / 1000.0)
    return int(s // 1)


def totals(events, day, now=None):
    """The day's counts; running feeds and sleeps fold `now`. Mirrors Core.totals exactly."""
    t = {"feeds": 0, "bottle_ml": 0, "breast_s": 0, "wet": 0, "dirty": 0, "sleeps": 0,
         "sleep_s": 0, "pumps": 0, "pump_ml": 0}
    for ev in on_day(events, day):
        data = ev.get("data") or {}
        kind = ev.get("type")
        if kind == "feed":
            t["feeds"] += 1
            t["bottle_ml"] += bottle_ml(ev)
            t["breast_s"] += breast_seconds(ev, now)
        elif kind == "diaper":
            if data.get("wet"):
                t["wet"] += 1
            if data.get("dirty"):
                t["dirty"] += 1
        elif kind == "sleep":
            t["sleeps"] += 1
            end = to_ms(now) if ev.get("end") is None else to_ms(ev["end"])
            start = to_ms(ev.get("time"))
            if end is not None and start is not None:
                t["sleep_s"] += max(0, int((end - start) // 1000))
        elif kind == "pump":
            t["pumps"] += 1
            t["pump_ml"] += (data.get("left_ml") or 0) + (data.get("right_ml") or 0)
    return t


def fmt_amount(ml, unit):
    """"22 ml" / "0.75 oz" — oz to the nearest quarter, as Core.fmtAmount."""
    if unit == "oz":
        q = round(ml / OZ_ML * 4) / 4
        return f"{q:g} oz"
    return f"{round(ml)} ml"


def fmt_time(iso):
    return str(iso)[11:16]


def fmt_day(iso):
    """"Wed 23 Sep" from the string's own date, never re-zoned."""
    s = str(iso)
    try:
        d = datetime(int(s[0:4]), int(s[5:7]), int(s[8:10]))
    except ValueError:
        return s[:10]
    return f"{DAYS[d.weekday()]} {d.day} {MONTHS[d.month - 1]}"


def since_text(ms):
    """"1 h 15 m", "45 m", "2 d 3 h": whole units, the trailing zero unit dropped."""
    t = max(0, int((ms or 0) // 60000))
    d, h, m = t // 1440, (t % 1440) // 60, t % 60
    if d > 0:
        return f"{d} d {h} h" if h else f"{d} d"
    if h > 0:
        return f"{h} h {m} m" if m else f"{h} h"
    return f"{m} m"


def days_between(a, b):
    return (datetime.strptime(b[:10], "%Y-%m-%d") - datetime.strptime(a[:10], "%Y-%m-%d")).days


def day_number(born, iso):
    """1-based day of life for the heading."""
    return days_between(born, local_date(iso)) + 1


def _label(s):
    return str(s).replace("_", " ")


def _num(v):
    """A number without a trailing .0, the way JS prints it."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def describe(ev, unit="ml", now=None):
    """The one-line summary: "5 min + 22 ml formula", "Wet + dirty · yellow", "Sleep 1 h 20 m"."""
    data = ev.get("data") or {}
    kind = ev.get("type")
    if kind == "feed":
        parts = []
        s = breast_seconds(ev, now)
        if s > 0:
            parts.append(f"{round(s / 60)} min")
        by_kind = {}
        for b in data.get("bottles") or []:
            by_kind[b["kind"]] = by_kind.get(b["kind"], 0) + (b.get("ml") or 0)
        for k, ml in by_kind.items():
            if ml:
                parts.append(f"{fmt_amount(ml, unit)} {_label(k)}")
        if parts:
            return " + ".join(parts)
        return "Feeding" if is_running(ev) else "Feed"
    if kind == "diaper":
        text = ("Wet + dirty" if data.get("wet") and data.get("dirty") else
                "Wet" if data.get("wet") else "Dirty" if data.get("dirty") else "Diaper")
        extra = [_label(x) for x in (data.get("color"), data.get("texture")) if x]
        if data.get("blowout"):
            extra.append("blowout")
        return f"{text} · {', '.join(extra)}" if extra else text
    if kind == "sleep":
        if is_running(ev):
            return f"Sleeping {since_text(to_ms(now) - to_ms(ev['time']))}"
        return f"Sleep {since_text(to_ms(ev['end']) - to_ms(ev['time']))}"
    if kind == "pump":
        ml = (data.get("left_ml") or 0) + (data.get("right_ml") or 0)
        if ml or data.get("left_ml") is not None or data.get("right_ml") is not None:
            return f"Pumped {fmt_amount(ml, unit)}"
        return "Pump"
    if kind == "growth":
        parts = []
        if data.get("weight_g") is not None:
            parts.append(f"Weight {_num(round(data['weight_g'] / 10) / 100)} kg")
        if data.get("length_cm") is not None:
            parts.append(f"Length {_num(data['length_cm'])} cm")
        if data.get("head_cm") is not None:
            parts.append(f"Head {_num(data['head_cm'])} cm")
        return " · ".join(parts) or "Growth"
    if kind == "health":
        parts = []
        if data.get("medicine"):
            parts.append(" ".join(x for x in (data["medicine"], data.get("dose")) if x))
        if data.get("temp_c") is not None:
            parts.append(f"{_num(data['temp_c'])} °C")
        if data.get("symptom"):
            parts.append(data["symptom"])
        return " · ".join(parts) or "Health"
    if kind == "note":
        return "Milestone" if data.get("milestone") else "Note"
    return str(kind or "Entry")


# -- the workbook ------------------------------------------------------------------------------

def _lock_path(target):
    return Path(target).with_name(f"~${Path(target).name}")


def workbook_path(output_folder):
    return Path(output_folder) / WORKBOOK


def is_locked(output_folder):
    """True when Excel has the workbook open (its `~$` lock file exists)."""
    return _lock_path(workbook_path(output_folder)).exists()


def _backup(target, backup_dir, keep):
    if not target.exists():
        return None
    backup_dir.mkdir(parents=True, exist_ok=True)
    # Microseconds, not seconds. Three regenerations inside one second produced one backup
    # filename and silently overwrote each other — the third instance of this bug in the whiskey
    # project. See SPEC.md §1.
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    dest = backup_dir / f"{target.stem}_{stamp}{target.suffix}"
    shutil.copy2(target, dest)
    old = sorted(backup_dir.glob(f"{target.stem}_*{target.suffix}"))
    for p in old[:-keep]:
        p.unlink(missing_ok=True)
    return dest


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, len(h) + 2))
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    return ws


def _minutes(seconds):
    return None if seconds is None else round(seconds / 60, 1)


def _yn(v):
    return "yes" if v else "no"


def _end_cell(ev):
    """The end clock, with its date in front only when it is not the start's date."""
    end = ev.get("end")
    if end is None:
        return None
    if local_date(end) == local_date(ev.get("time")):
        return fmt_time(end)
    return f"{local_date(end)} {fmt_time(end)}"


def _sleep_minutes(ev, now):
    end = to_ms(now) if ev.get("end") is None else to_ms(ev["end"])
    start = to_ms(ev.get("time"))
    if end is None or start is None:
        return None
    return round(max(0.0, (end - start) / 1000) / 60, 1)


def _common(ev):
    return [local_date(ev.get("time")), fmt_time(ev.get("time"))]


def _feed_row(ev, now):
    d = ev["data"]
    br = d.get("breast") or {}
    by_kind = {}
    for b in d.get("bottles") or []:
        by_kind[b["kind"]] = by_kind.get(b["kind"], 0) + (b.get("ml") or 0)
    bottle = sum(by_kind.values())
    return _common(ev) + [
        _end_cell(ev), ev.get("logged_by"),
        _minutes(breast_seconds(ev, now)) if breast_seconds(ev, now) else None,
        _minutes(br.get("left_s")), _minutes(br.get("right_s")),
        bottle or None, by_kind.get("formula"), by_kind.get("breast_milk"),
        d.get("made_ml"), d.get("leftover_ml"), ev.get("note") or None, ev["event_id"]]


def _diaper_row(ev):
    d = ev["data"]
    return _common(ev) + [ev.get("logged_by"), _yn(d.get("wet")), _yn(d.get("dirty")),
                          d.get("color"), d.get("texture"), d.get("size"), _yn(d.get("rash")),
                          _yn(d.get("blowout")), ev.get("note") or None, ev["event_id"]]


def _sleep_row(ev, now):
    d = ev["data"]
    return _common(ev) + [_end_cell(ev), _sleep_minutes(ev, now), d.get("where"),
                          ev.get("logged_by"), ev.get("note") or None, ev["event_id"]]


def _pump_row(ev):
    d = ev["data"]
    total = (d.get("left_ml") or 0) + (d.get("right_ml") or 0)
    has = d.get("left_ml") is not None or d.get("right_ml") is not None
    return _common(ev) + [ev.get("logged_by"), d.get("left_ml"), d.get("right_ml"),
                          total if has else None, d.get("minutes"), ev.get("note") or None,
                          ev["event_id"]]


def _growth_row(ev):
    d = ev["data"]
    return _common(ev) + [ev.get("logged_by"), d.get("weight_g"), d.get("length_cm"),
                          d.get("head_cm"), ev.get("note") or None, ev["event_id"]]


def _health_row(ev):
    d = ev["data"]
    return _common(ev) + [ev.get("logged_by"), d.get("temp_c"), d.get("medicine"), d.get("dose"),
                          d.get("symptom"), ev.get("note") or None, ev["event_id"]]


def _note_row(ev):
    return _common(ev) + [ev.get("logged_by"), _yn(ev["data"].get("milestone")),
                          ev.get("note") or None, ev["event_id"]]


def daily_rows(events, d1, d2, now=None):
    """One dict per local date from d1 to d2 inclusive with the Daily columns; zeros on empty
    days, minutes rounded to one place."""
    out = []
    try:
        day = datetime.strptime(d1, "%Y-%m-%d")
        last = datetime.strptime(d2, "%Y-%m-%d")
    except (TypeError, ValueError):
        return out
    while day <= last:
        date = day.strftime("%Y-%m-%d")
        t = totals(events, date, now)
        out.append({"date": date, "feeds": t["feeds"], "bottle_ml": t["bottle_ml"],
                    "breast_min": round(t["breast_s"] / 60, 1), "wet": t["wet"],
                    "dirty": t["dirty"], "sleeps": t["sleeps"],
                    "sleep_min": round(t["sleep_s"] / 60, 1), "pumps": t["pumps"],
                    "pump_ml": t["pump_ml"]})
        day += timedelta(days=1)
    return out


def _warn_unreadable(journal):
    """A placeholder or half-synced file is skipped, not fatal — but it is said out loud, because
    a workbook missing one feed looks exactly like a workbook with every feed."""
    for name in journal.stats().get("unreadable") or []:
        log.warning("rollup: skipping unreadable journal file %s", name)


def build(journal, output_folder, backup_dir, keep=30):
    """Regenerate `Baby Log.xlsx` in the output folder. Returns a summary dict."""
    target = workbook_path(output_folder)
    if _lock_path(target).exists():
        raise RuntimeError(
            f"{target.name} is open in Excel ({_lock_path(target).name} present). "
            "Close it and run again — writing now would create a OneDrive conflict copy.")
    backup_dir = Path(backup_dir)
    now = datetime.now().astimezone()

    _warn_unreadable(journal)
    events = journal.events()                       # live, oldest first (§3.4)
    by_type = {t: [] for t in store_mod.TYPES}
    for ev in events:
        by_type.setdefault(ev.get("type"), []).append(ev)

    wb = Workbook()
    counts = {}
    for title, kind, cols, make in (
            ("Feeds", "feed", FEED_COLUMNS, lambda e: _feed_row(e, now)),
            ("Diapers", "diaper", DIAPER_COLUMNS, _diaper_row),
            ("Sleep", "sleep", SLEEP_COLUMNS, lambda e: _sleep_row(e, now)),
            ("Pumping", "pump", PUMP_COLUMNS, _pump_row),
            ("Growth", "growth", GROWTH_COLUMNS, _growth_row),
            ("Health", "health", HEALTH_COLUMNS, _health_row),
            ("Notes", "note", NOTE_COLUMNS, _note_row)):
        rows = [make(e) for e in by_type.get(kind, [])]
        _sheet(wb, title, cols, rows, {"note": 40, "event_id": 30, "logged_by": 12, "date": 12,
                                       "medicine": 16, "symptom": 24, "where": 16})
        counts[title.lower()] = len(rows)

    first = local_date(events[0]["time"]) if events else None
    drows = daily_rows(events, first, today_local(now), now) if first else []
    _sheet(wb, "Daily", DAILY_COLUMNS, [[r[c] for c in DAILY_COLUMNS] for r in drows],
           {"date": 12})
    counts["daily"] = len(drows)

    _sheet(wb, "About", ["field", "value"], [
        ["generated", now.isoformat(timespec="seconds")],
        ["source", "journal at " + str(journal.root)],
        ["WARNING", "Generated file. Rewritten from the journal whenever it changes — "
                    "do not hand-edit, changes will be lost. Correct entries in the app."],
        ["minutes", "*_min columns are seconds / 60 rounded to one place"],
        ["daily", "one row per local date from the first entry to today; zeros when empty"],
    ], {"field": 12, "value": 90})

    _backup(target, backup_dir, keep)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(target.parent), prefix=".tmp-", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, target)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
    finally:
        wb.close()

    return {"path": str(target), **counts, "bytes": target.stat().st_size}


# -- the day sheet -----------------------------------------------------------------------------

def _css():
    """static/print.css inlined; the sheet must print with no network at all."""
    try:
        return PRINT_CSS.read_text(encoding="utf-8")
    except OSError:
        log.warning("print.css not found at %s; using a bare style", PRINT_CSS)
        return "@page { size: Letter; margin: 0.5in } body { color: #000; background: #fff }"


def _esc(v):
    return html.escape("" if v is None else str(v), quote=True)


def _breast_cell(ev, now):
    br = (ev.get("data") or {}).get("breast") or {}
    s = breast_seconds(ev, now)
    if not s and br.get("total_s") is None and br.get("left_s") is None and br.get("right_s") is None:
        return ""
    text = f"{round(s / 60)} min"
    sides = []
    if br.get("left_s") is not None:
        sides.append(f"L {round(br['left_s'] / 60)}")
    if br.get("right_s") is not None:
        sides.append(f"R {round(br['right_s'] / 60)}")
    if sides:
        text += f" ({'/'.join(sides)})"
    if br.get("approx"):
        text = "~" + text
    if is_running(ev):
        text += " …"
    return text


def _bottle_cell(ev, unit):
    parts = []
    for b in (ev.get("data") or {}).get("bottles") or []:
        parts.append(f"{fmt_amount(b.get('ml') or 0, unit)} {_label(b['kind'])}")
    return " + ".join(parts)


def _made_cell(ev, unit):
    d = ev.get("data") or {}
    made, left = d.get("made_ml"), d.get("leftover_ml")
    if made is None and left is None:
        return ""
    a = fmt_amount(made, unit) if made is not None else "–"
    b = fmt_amount(left, unit) if left is not None else "–"
    return f"{a} / {b}"


def _feed_table(feeds, unit, now):
    head = ("<tr><th>Time</th><th>Breast</th><th>Bottle</th><th>Made / Leftover</th>"
            "<th>By</th><th>Note</th></tr>")
    rows = []
    for ev in feeds:
        rows.append(
            f"<tr><td class=\"time\">{_esc(fmt_time(ev['time']))}</td>"
            f"<td>{_esc(_breast_cell(ev, now))}</td>"
            f"<td>{_esc(_bottle_cell(ev, unit))}</td>"
            f"<td>{_esc(_made_cell(ev, unit))}</td>"
            f"<td>{_esc(ev.get('logged_by'))}</td>"
            f"<td>{_esc(ev.get('note'))}</td></tr>")
    if not rows:
        rows.append("<tr><td colspan=\"6\" class=\"empty\">No feeds</td></tr>")
    return f"<table class=\"feeds\">{head}{''.join(rows)}</table>"


def _diaper_table(diapers):
    head = ("<tr><th>Time</th><th>Wet</th><th>Dirty</th><th>Colour / texture</th>"
            "<th>By</th><th>Note</th></tr>")
    rows = []
    for ev in diapers:
        d = ev.get("data") or {}
        extra = [_label(x) for x in (d.get("color"), d.get("texture")) if x]
        if d.get("size"):
            extra.append(_label(d["size"]))
        if d.get("rash"):
            extra.append("rash")
        if d.get("blowout"):
            extra.append("blowout")
        rows.append(
            f"<tr><td class=\"time\">{_esc(fmt_time(ev['time']))}</td>"
            f"<td class=\"mark\">{'✓' if d.get('wet') else ''}</td>"
            f"<td class=\"mark\">{'✓' if d.get('dirty') else ''}</td>"
            f"<td>{_esc(', '.join(extra))}</td>"
            f"<td>{_esc(ev.get('logged_by'))}</td>"
            f"<td>{_esc(ev.get('note'))}</td></tr>")
    if not rows:
        rows.append("<tr><td colspan=\"6\" class=\"empty\">No diapers</td></tr>")
    return f"<table class=\"diapers\">{head}{''.join(rows)}</table>"


def _other_table(others, unit, now):
    head = "<tr><th>Time</th><th>Type</th><th>Details</th><th>By</th><th>Note</th></tr>"
    rows = []
    for ev in others:
        rows.append(
            f"<tr><td class=\"time\">{_esc(fmt_time(ev['time']))}</td>"
            f"<td>{_esc(str(ev.get('type', '')).capitalize())}</td>"
            f"<td>{_esc(describe(ev, unit, now))}</td>"
            f"<td>{_esc(ev.get('logged_by'))}</td>"
            f"<td>{_esc(ev.get('note'))}</td></tr>")
    return f"<table class=\"other\">{head}{''.join(rows)}</table>"


def footer_text(t, unit, targets=None):
    """`8 feeds · 135 ml bottle · 57 min breast · 6 wet · 4 dirty`, targets beside a count when
    the pediatrician gave one."""
    targets = targets or {}

    def with_target(n, key):
        goal = targets.get(key)
        return f"{n} (target {_num(goal)})" if goal is not None else str(n)

    parts = [f"{with_target(t['feeds'], 'feeds_per_day')} feeds",
             f"{fmt_amount(t['bottle_ml'], unit)} bottle",
             f"{round(t['breast_s'] / 60)} min breast",
             f"{with_target(t['wet'], 'wet_per_day')} wet",
             f"{with_target(t['dirty'], 'dirty_per_day')} dirty"]
    if t["sleeps"]:
        parts.append(f"{t['sleeps']} sleeps · {since_text(t['sleep_s'] * 1000)}")
    if t["pumps"]:
        parts.append(f"pumped {fmt_amount(t['pump_ml'], unit)}")
    return " · ".join(parts)


def _page(title, body):
    return ("<!DOCTYPE html>\n<html lang=\"en\"><head><meta charset=\"utf-8\">\n"
            f"<title>{_esc(title)}</title>\n<style>\n{_css()}\n</style></head>\n"
            f"<body>\n{body}\n</body></html>\n")


def day_sheet_html(child, events, date, settings, now=None):
    """The one renderer for a day: served at /print/day/<date> and saved to Day sheets/."""
    child = child or {}
    settings = settings or {}
    unit = settings.get("units") or "ml"
    name = child.get("name") or "Baby"
    born = child.get("born")
    day_n = day_number(born, date) if born else None
    heading = f"{name} — Feeding & Diapering — {fmt_day(date)}"
    if day_n is not None:
        heading += f" (day {day_n})"
    todays = sorted(on_day(events, date), key=lambda e: (to_ms(e.get("time")) or 0, e.get("event_id")))
    feeds = [e for e in todays if e.get("type") == "feed"]
    diapers = [e for e in todays if e.get("type") == "diaper"]
    others = [e for e in todays if e.get("type") not in ("feed", "diaper")]
    t = totals(todays, date, now)

    body = [f"<h1>{_esc(heading)}</h1>",
            "<div class=\"columns\">",
            f"<div class=\"column\"><h2>Feeding</h2>{_feed_table(feeds, unit, now)}</div>",
            f"<div class=\"column\"><h2>Diapers</h2>{_diaper_table(diapers)}</div>",
            "</div>"]
    if others:
        body.append(f"<div class=\"other\"><h2>Other</h2>{_other_table(others, unit, now)}</div>")
    body.append(f"<p class=\"footer\">{_esc(footer_text(t, unit, child.get('targets')))}</p>")
    return _page(f"{name} — {date}", "\n".join(body))


def range_html(child, rows, d1, d2, settings):
    """The daily-totals summary for a date range: one table with the Daily columns."""
    child = child or {}
    settings = settings or {}
    unit = settings.get("units") or "ml"
    name = child.get("name") or "Baby"
    heading = f"{name} — Daily totals — {fmt_day(d1)} to {fmt_day(d2)}"
    num = ' class="num"'
    head = "".join(f"<th{num if c != 'date' else ''}>{_esc(c)}</th>" for c in DAILY_COLUMNS)
    body_rows = []
    for r in rows:
        cells = []
        for c in DAILY_COLUMNS:
            v = r.get(c)
            if c in ("bottle_ml", "pump_ml") and unit == "oz":
                v = fmt_amount(v or 0, unit)
            text = _num(v) if isinstance(v, (int, float)) else v
            cells.append(f"<td{num if c != 'date' else ''}>{_esc(text)}</td>")
        body_rows.append(f"<tr>{''.join(cells)}</tr>")
    if not body_rows:
        body_rows.append(f"<tr><td colspan=\"{len(DAILY_COLUMNS)}\" class=\"empty\">No entries</td></tr>")
    table = f"<table class=\"range\"><tr>{head}</tr>{''.join(body_rows)}</table>"
    return _page(f"{name} — {d1} to {d2}", f"<h1>{_esc(heading)}</h1>\n{table}")


def write_day_sheet(journal, output_folder, date, settings, child=None):
    """Save the day sheet as `Day sheets/<date>.html` in the output folder; returns the path."""
    if child is None:
        kids = journal.children()
        child = kids[0] if kids else None
    text = day_sheet_html(child, journal.events(), date, settings)
    folder = Path(output_folder) / DAY_SHEETS
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{date}.html"
    fd, tmp = tempfile.mkstemp(dir=str(folder), prefix=".tmp-", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(text)
        os.replace(tmp, target)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
    return target


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="Regenerate Baby Log.xlsx from the journal")
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--root", help="journal root (also BABY_LOG_APP_FOLDER)")
    ap.add_argument("--out", help="override the output folder (also BABY_LOG_OUTPUT_FOLDER)")
    a = ap.parse_args(argv)

    with open(a.config, encoding="utf-8") as fh:
        cfg = yaml.safe_load(fh) or {}
    paths = cfg.get("paths") or {}
    j = store_mod.Journal(store_mod.resolve_root(cfg, a.root)).ensure()
    out = Path(a.out or os.environ.get("BABY_LOG_OUTPUT_FOLDER") or paths["output_folder"])
    backups = Path(paths.get("backups") or Path(paths.get("local_data") or "./data") / "backups")
    summary = build(j, out, backups, int(paths.get("backup_keep", 30)))
    for k, v in summary.items():
        print(f"  {k:<11} {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
